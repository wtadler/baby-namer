#!/usr/bin/env python3
"""
Process ONS England & Wales baby names to data/england-wales/names.csv.

The ONS provides a single Excel workbook with two tables (girls and boys),
each in wide format: one row per name, alternating Rank/Count columns per year.

Download the workbook from:
  https://www.ons.gov.uk/peoplepopulationandcommunity/birthsdeathsandmarriages/
  livebirths/datasets/babynamesinenglandandwalesfrom1996

Save as:
  data/england-wales/raw/babynames1996to2024.xlsx

Requires: pip install openpyxl
"""
import csv
import os
import sys
import re

OUT      = "data/england-wales/names.csv"
RAW_FILE = "data/england-wales/raw/babynames1996to2024.xlsx"

# Also support separate boys/girls files (legacy fallback)
BOYS_FILE  = "data/england-wales/raw/boys.xlsx"
GIRLS_FILE = "data/england-wales/raw/girls.xlsx"

try:
    import openpyxl
except ImportError:
    print("Missing dependency. Run: pip install openpyxl")
    sys.exit(1)


def read_wide_table(ws, sex_label):
    """Read a wide-format ONS sheet (Name, YYYY Rank, YYYY Count, ...) into rows."""
    rows_out = []
    all_rows = list(ws.iter_rows(values_only=True))

    # Find header row: first row where first cell is 'Name'
    header_row_idx = None
    for i, row in enumerate(all_rows):
        if row and str(row[0]).strip().lower() == 'name':
            header_row_idx = i
            break

    if header_row_idx is None:
        print(f"  Warning: no header row found in sheet, skipping")
        return rows_out

    header = all_rows[header_row_idx]

    # Parse year->count_col mapping from header
    # Header looks like: Name, 2024 Rank, 2024 Count, 2023 Rank, 2023 Count, ...
    year_count_cols = {}  # year -> col index
    for col_idx, cell in enumerate(header):
        if not cell:
            continue
        m = re.match(r'(\d{4})\s+count', str(cell).strip(), re.IGNORECASE)
        if m:
            year_count_cols[int(m.group(1))] = col_idx

    if not year_count_cols:
        print(f"  Warning: no year/count columns found, skipping")
        return rows_out

    for row in all_rows[header_row_idx + 1:]:
        if not row or not row[0]:
            continue
        name_val = row[0]
        if not isinstance(name_val, str):
            continue
        name = name_val.strip().title()
        if not name:
            continue

        for year, count_col in year_count_cols.items():
            if count_col >= len(row):
                continue
            count_val = row[count_col]
            if count_val is None or count_val == '' or str(count_val).startswith('['):
                continue
            try:
                count = int(count_val)
            except (ValueError, TypeError):
                continue
            if count <= 0:
                continue
            rows_out.append((year, sex_label, name, count))

    return rows_out


def read_ons_workbook_wide(path):
    """Read the combined ONS workbook with Table_1 (girls) and Table_2 (boys)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    all_rows = []

    # Determine which sheets to use
    sheet_map = {}
    for name in wb.sheetnames:
        n = name.strip().lower()
        if 'table_1' in n or 'table 1' in n or ('girl' in n and 'table' in n):
            sheet_map['Girl'] = name
        elif 'table_2' in n or 'table 2' in n or ('boy' in n and 'table' in n):
            sheet_map['Boy'] = name

    if not sheet_map:
        print("  Warning: could not identify Girl/Boy tables, trying Table_1/Table_2")
        if 'Table_1' in wb.sheetnames:
            sheet_map['Girl'] = 'Table_1'
        if 'Table_2' in wb.sheetnames:
            sheet_map['Boy'] = 'Table_2'

    for sex_label, sheet_name in sheet_map.items():
        ws = wb[sheet_name]
        rows = read_wide_table(ws, sex_label)
        years = {r[0] for r in rows}
        print(f"  {sex_label}: {len(rows):,} rows across {len(years)} years ({min(years)}–{max(years)})")
        all_rows.extend(rows)

    wb.close()
    return all_rows


def read_ons_workbook_by_year(path, sex_label):
    """Fallback: read a workbook where each sheet is a year (old format)."""
    rows = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        try:
            year = int(sheet_name.strip())
        except ValueError:
            continue

        ws = wb[sheet_name]
        data_rows = list(ws.iter_rows(values_only=True))

        header_row_idx = None
        for i, row in enumerate(data_rows):
            cells = [str(c).lower() if c else '' for c in row]
            if any('name' in c for c in cells) and any('count' in c or 'number' in c for c in cells):
                header_row_idx = i
                break

        if header_row_idx is None:
            continue

        headers = [str(c).lower().strip() if c else '' for c in data_rows[header_row_idx]]
        try:
            name_col  = next(i for i, h in enumerate(headers) if 'name' in h)
            count_col = next(i for i, h in enumerate(headers) if 'count' in h or 'number' in h)
        except StopIteration:
            continue

        for row in data_rows[header_row_idx + 1:]:
            if not row or len(row) <= max(name_col, count_col):
                continue
            name_val  = row[name_col]
            count_val = row[count_col]
            if not name_val or not count_val:
                continue
            name = str(name_val).strip().title()
            try:
                count = int(count_val)
            except (ValueError, TypeError):
                continue
            if count <= 0:
                continue
            rows.append((year, sex_label, name, count))

    wb.close()
    return rows


# ── Main ─────────────────────────────────────────────────────────────────────

os.makedirs("data/england-wales", exist_ok=True)

if os.path.exists(RAW_FILE):
    print(f"Processing combined workbook: {RAW_FILE}")
    all_rows = read_ons_workbook_wide(RAW_FILE)

elif os.path.exists(BOYS_FILE) and os.path.exists(GIRLS_FILE):
    print(f"Processing separate boys/girls workbooks ...")
    print(f"  Boys: {BOYS_FILE}")
    boys = read_ons_workbook_by_year(BOYS_FILE, 'Boy')
    print(f"    {len(boys):,} rows across {len({r[0] for r in boys})} years")
    print(f"  Girls: {GIRLS_FILE}")
    girls = read_ons_workbook_by_year(GIRLS_FILE, 'Girl')
    print(f"    {len(girls):,} rows across {len({r[0] for r in girls})} years")
    all_rows = boys + girls

else:
    print("ONS data file not found. Please download it:\n")
    print("1. Go to https://www.ons.gov.uk/peoplepopulationandcommunity/birthsdeathsandmarriages/livebirths/datasets/babynamesinenglandandwalesfrom1996")
    print("2. Download the Excel workbook (1996 to current)")
    print("3. Save as:")
    print(f"     {RAW_FILE}")
    print("4. Run: python3 process_england_wales.py")
    sys.exit(1)

if not all_rows:
    print("ERROR: no data rows extracted. Check the workbook format.")
    sys.exit(1)

all_rows = sorted(all_rows, key=lambda r: (r[0], r[1], r[2]))

with open(OUT, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(['Year', 'Sex', 'Name', 'Count'])
    w.writerows(all_rows)

years = sorted({r[0] for r in all_rows})
print(f"\nWritten {len(all_rows):,} rows to {OUT}")
print(f"Year range: {min(years)}–{max(years)}, unique names: {len({r[2] for r in all_rows}):,}")
